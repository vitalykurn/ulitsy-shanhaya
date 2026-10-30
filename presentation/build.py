# -*- coding: utf-8 -*-
"""Build the 16:9 pptx from presentation/content.py."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

import content

ROOT = Path(__file__).resolve().parent
IMG = ROOT.parent / "site" / "img"
OUT = ROOT / "ulitsy-shanhaya.pptx"

PAPER = RGBColor(0xF3, 0xEE, 0xE4)
INK = RGBColor(0x24, 0x1C, 0x16)
BRONZE = RGBColor(0x8A, 0x4E, 0x2B)
MUTED = RGBColor(0x6D, 0x5E, 0x52)
WHITE = RGBColor(0xFF, 0xFA, 0xF3)


def set_run_font(run, name, size, color, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = rpr.makeelement(qn(f"a:{tag}"), {})
            rpr.append(el)
        el.set("typeface", name)


def add_text(slide, left, top, width, height, lines, size=20, color=INK, font="Calibri", bold=False, align=PP_ALIGN.LEFT, zh=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(8)
        run = paragraph.add_run()
        run.text = line
        set_run_font(run, font, size, color, bold=bold)
    if zh:
        paragraph = frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(12)
        run = paragraph.add_run()
        run.text = zh
        set_run_font(run, "Microsoft YaHei", 15, MUTED)
    return box


def paint(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), Inches(7.5))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BRONZE
    bar.line.fill.background()
    foot = slide.shapes.add_textbox(Inches(0.45), Inches(7.08), Inches(12.4), Inches(0.32))
    paragraph = foot.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = "Вера Курносенко  ·  Шанхай  ·  ulitsy-shanhaya.ru"
    set_run_font(run, "Calibri", 12, MUTED)


def add_picture(slide, name, left, top, max_w, max_h):
    path = IMG / name
    if not path.exists():
        return None
    with Image.open(path) as image:
        ratio = image.width / image.height
    width = max_w
    height = width / ratio
    if height > max_h:
        height = max_h
        width = height * ratio
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))
    return top + height


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "История России на улицах Шанхая"
    prs.core_properties.author = "Курносенко Вера Витальевна"
    prs.core_properties.subject = "Карта для китайских одноклассников"
    blank = prs.slide_layouts[6]

    total = len(content.SLIDES)
    for index, slide_data in enumerate(content.SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        paint(slide)
        add_text(slide, 0.45, 0.28, 8, 0.32, [slide_data.get("kicker", "")], size=14, color=BRONZE, bold=True)
        title_lines = slide_data["title"].split("\n")
        title_size = 32 if len(title_lines) > 1 or len(slide_data["title"]) > 32 else 34
        add_text(slide, 0.42, 0.58, 7.6 if slide_data.get("kind") == "photo" else 12.2, 1.55, title_lines, size=title_size, color=INK, font="Georgia", bold=True)
        has_photo = slide_data.get("kind") == "photo" and slide_data.get("image")
        picture = slide_data.get("picture") or {}
        text_width = picture.get("text_width", 6.7 if has_photo else 12.2)
        lines = slide_data["lines"]
        body_size = 18 if len(lines) >= 5 else 20 if len(lines) == 4 else 22
        body_top = 2.25
        body_height = 4.55
        add_text(
            slide, 0.48, body_top, text_width, body_height, lines,
            size=body_size, color=INK, zh=slide_data.get("zh"),
        )
        if has_photo:
            bottom = add_picture(
                slide,
                slide_data["image"],
                picture.get("left", 7.55),
                picture.get("top", 1.15),
                picture.get("max_w", 5.3),
                picture.get("max_h", 5.15),
            )
            caption = slide_data.get("caption")
            if caption and bottom is not None and bottom < 6.35:
                add_text(slide, 7.55, bottom + 0.08, 5.2, 0.7, [caption], size=12, color=MUTED)
        num = slide.shapes.add_textbox(Inches(12.15), Inches(7.05), Inches(0.95), Inches(0.32))
        run = num.text_frame.paragraphs[0].add_run()
        run.text = f"{index}/{total}"
        set_run_font(run, "Calibri", 12, MUTED)
        num.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT

    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    lines = []
    for slide_data in content.SLIDES:
        lines.append(slide_data.get("kicker", ""))
        lines.append(slide_data["title"].replace("\n", " "))
        lines.extend(slide_data["lines"])
        if slide_data.get("caption"):
            lines.append(slide_data["caption"])
    texts = ROOT.parent / "texts"
    texts.mkdir(exist_ok=True)
    (texts / "slides-ru.txt").write_text("\n\n".join(lines), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    build()
